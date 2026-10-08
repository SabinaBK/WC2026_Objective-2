"""
build_datasets.py
-----------------
Builds the two modelling datasets from WC2026_data.xlsx:

    lr21.csv  - Linear Regression 2.1, 104 rows (one per match),  Y = goal_diff
    lr22.csv  - Linear Regression 2.2, 208 rows (two per match),  Y = goals

Every variable uses only information available BEFORE kick-off:
  * recent form uses matches dated strictly before the match date
  * in-tournament xG uses only earlier 2026 World Cup matches
  * FIFA points come from the last ranking released before 11 June 2026
  * player data is from the completed 2025-26 club season

Usage:
    python build_datasets.py WC2026_data.xlsx

Requirements: pandas, numpy, openpyxl
"""
import sys
import numpy as np
import pandas as pd
from openpyxl import load_workbook

# ----------------------------------------------------------------- settings
WORKBOOK = sys.argv[1] if len(sys.argv) > 1 else "WC2026_data.xlsx"
TOURNAMENT_START = pd.Timestamp("2026-06-11")
HOSTS = {"USA", "Canada", "Mexico"}      
N_FORM = 10            # last N competitive matches
DECAY = 0.9            # weight of the match i places back = DECAY ** i
K_BLEND = 3            # prior counts as K_BLEND matches in the blended xG variables
MIN_MINUTES = 900      # minimum league minutes for attackers in attack_p90
DEFAULT_TIER = 0.5     # weight for leagues not listed in League_Tiers

warnings = []
def warn(msg):
    warnings.append(msg)

SHEETS = pd.ExcelFile(WORKBOOK).sheet_names

def read_tab(name, ncols, names, key):
    """Read a workbook tab by column position, drop empty / note rows.
    A sheet that does not exist yet returns an empty table (pilot mode)."""
    if name not in SHEETS:
        warn(f"Sheet '{name}' not found yet - variables that need it will be blank.")
        return pd.DataFrame(columns=names)
    df = pd.read_excel(WORKBOOK, sheet_name=name, header=0)
    df = df.iloc[:, :ncols]
    if df.shape[1] < ncols:
        sys.exit(f"Sheet '{name}' needs {ncols} columns in this order: {names}")
    df.columns = names
    df = df[df[key].notna()].copy()
    return df

def clean_name(s):
    return s.strip() if isinstance(s, str) else s

# ------------------------------------------------------------------- load raw data
split = read_tab("Split", 4, ["match_no", "stage", "Set_80_20", "Set_70_30"], "match_no")
split["match_no"] = pd.to_numeric(split["match_no"], errors="coerce")
split = split.dropna(subset=["match_no"]).astype({"match_no": int})
split["match_no"] = split["match_no"].astype(float)

m = read_tab("Matches", 8, ["match_no", "date", "stage", "team_A", "team_B",
                            "goals_A", "goals_B", "shootout"], "match_no")
m["date"] = pd.to_datetime(m["date"], errors="coerce")
for c in ("team_A", "team_B"):
    m[c] = m[c].map(clean_name)
m = m.merge(split[["match_no", "Set_80_20", "Set_70_30"]], on="match_no", how="left")
m["match_no"] = pd.to_numeric(m["match_no"], errors="coerce")
m_filled = m.dropna(subset=["date", "team_A", "team_B"]).copy()
if len(m_filled) < len(m):
    warn(f"Matches: {len(m) - len(m_filled)} rows have no date/teams yet - they are skipped.")

fp = read_tab("FIFA_Points", 4, ["team", "points", "release", "qualified"], "team")
fp["team"] = fp["team"].map(clean_name)
P = dict(zip(fp["team"], pd.to_numeric(fp["points"], errors="coerce")))
qual = fp[fp["qualified"].astype(str).str.upper().str.strip() == "Y"]
P_AVG = pd.to_numeric(qual["points"], errors="coerce").mean()
if np.isnan(P_AVG) and len(P):
    P_AVG = np.nanmean(list(P.values()))
    warn("FIFA_Points: no team marked qualified = Y; P_avg uses all teams listed.")

rm = read_tab("Recent_Matches", 6, ["team", "date", "opponent", "competition", "gf", "ga"], "date")
rm["date"] = pd.to_datetime(rm["date"], errors="coerce")
rm["team"] = rm["team"].map(clean_name)
rm["opponent"] = rm["opponent"].map(clean_name)
rm = rm[rm["date"] < TOURNAMENT_START]          # World Cup matches come from Matches

ms = read_tab("Match_Stats", 3, ["match_no", "team_ignored", "xg"], "match_no")
ms["side"] = ms.groupby("match_no").cumcount().map({0: "A", 1: "B"})   # 1st row = Team A

pl = read_tab("Players", 10, ["team", "player", "pos", "club", "league", "minutes",
                              "npg", "assists", "ucl_min", "ucl_yn"], "player")
pl["team"] = pl["team"].map(clean_name)
tiers = read_tab("League_Tiers", 2, ["league", "weight"], "weight")
TIER = dict(zip(tiers["league"].map(clean_name), tiers["weight"]))
pl["tier"] = pl["league"].map(clean_name).map(TIER).fillna(DEFAULT_TIER)
for c in ("minutes", "npg", "assists", "ucl_min"):
    pl[c] = pd.to_numeric(pl[c], errors="coerce").fillna(0)
pl["pos"] = pl["pos"].astype(str).str.upper().str.strip()

h2h = read_tab("H2H", 6, ["match_no", "a", "b", "meetings", "a_goals", "b_goals"], "match_no")

_wb = load_workbook(WORKBOOK, data_only=True, read_only=True)
XG22 = _wb["Constants"]["B3"].value if "Constants" in _wb.sheetnames else None
if XG22 is None:
    XG22 = np.nan
    warn("Constants!B3 (2022 World Cup average per team) is empty - 2.2 xg_for / opp_xga will be blank.")

# ------------------------------------------------------------------- checks
teams = sorted(set(m_filled["team_A"]) | set(m_filled["team_B"]))
for t in teams:
    if t not in P or pd.isna(P.get(t)):
        warn(f"FIFA_Points: no points for '{t}' (check spelling).")
    n_recent = (rm["team"] == t).sum()
    if n_recent < N_FORM:
        warn(f"Recent_Matches: '{t}' has {n_recent} matches (expected {N_FORM}).")
    if len(pl) and (pl["team"] == t).sum() == 0:
        warn(f"Players: no players entered for '{t}'.")
for opp in sorted(set(rm["opponent"].dropna())):
    if opp not in P or pd.isna(P.get(opp)):
        warn(f"FIFA_Points: no points for recent opponent '{opp}' - average points used instead.")

def p_of(team):
    v = P.get(team, np.nan)
    return P_AVG if pd.isna(v) else v

# ------------------------------------------------------------------- match history
hist_rows = [dict(team=r.team, date=r.date, gf=r.gf, ga=r.ga, opp=r.opponent) for r in rm.itertuples()]
played = m_filled.dropna(subset=["goals_A", "goals_B"])
for r in played.itertuples():
    hist_rows.append(dict(team=r.team_A, date=r.date, gf=r.goals_A, ga=r.goals_B, opp=r.team_B))
    hist_rows.append(dict(team=r.team_B, date=r.date, gf=r.goals_B, ga=r.goals_A, opp=r.team_A))
hist = pd.DataFrame(hist_rows)
hist["gf"] = pd.to_numeric(hist["gf"], errors="coerce")
hist["ga"] = pd.to_numeric(hist["ga"], errors="coerce")
hist["p_opp"] = hist["opp"].map(p_of)

def form(team, before):
    """Opponent-adjusted scoring and conceding form from the last N competitive matches."""
    h = hist[(hist["team"] == team) & (hist["date"] < before)].dropna(subset=["gf", "ga"])
    h = h.sort_values("date", ascending=False).head(N_FORM)
    if h.empty:
        return np.nan, np.nan
    if np.isnan(P_AVG):
        return np.nan, np.nan
    w = DECAY ** np.arange(len(h))
    gf, ga, po = h["gf"].to_numpy(float), h["ga"].to_numpy(float), h["p_opp"].to_numpy(float)
    scored = np.sum(w * gf * (po / P_AVG)) / w.sum()
    conceded = np.sum(w * ga * (P_AVG / po)) / w.sum()
    return scored, conceded

# ------------------------------------------------------------------- in-tournament xG
xg = ms.merge(m_filled[["match_no", "date", "team_A", "team_B"]], on="match_no", how="inner")
xg["team"] = np.where(xg["side"] == "A", xg["team_A"], xg["team_B"])
xg["xg"] = pd.to_numeric(xg["xg"], errors="coerce")
other = xg[["match_no", "side", "xg"]].copy()
other["side"] = other["side"].map({"A": "B", "B": "A"})
xg = xg.merge(other.rename(columns={"xg": "xga"}), on=["match_no", "side"], how="left")
xg = xg.rename(columns={"xg": "xgf"})[["team", "date", "match_no", "xgf", "xga"]]

def tournament_xg(team, before):
    h = xg[(xg["team"] == team) & (xg["date"] < before)]
    missing = h["xgf"].isna() | h["xga"].isna()
    if missing.any():
        warn(f"Match_Stats: xG missing for '{team}' in match(es) {sorted(h.loc[missing, 'match_no'].tolist())}.")
    return h[~missing]

def blend(n, mean, prior):
    if np.isnan(prior):
        return np.nan
    if n == 0:
        return prior
    return (n * mean + K_BLEND * prior) / (n + K_BLEND)

# ------------------------------------------------------------------- player variables
def team_players(t):
    return pl[pl["team"] == t]

def ucl_minutes(t):
    p = team_players(t)
    return p["ucl_min"].sum() if len(p) else np.nan

def attack_output(t):
    p = team_players(t)
    return ((p["npg"] + p["assists"]) * p["tier"]).sum() if len(p) else np.nan

def attack_p90(t):
    p = team_players(t)
    if not len(p):
        return np.nan
    p = p[p["minutes"] >= MIN_MINUTES]
    fw = p[p["pos"] == "FW"].sort_values("minutes", ascending=False)
    mf = p[p["pos"] == "MF"].sort_values("minutes", ascending=False)
    top = pd.concat([fw, mf]).head(3)      # forwards first, then midfielders
    if top.empty:
        warn(f"Players: '{t}' has no FW/MF with >= {MIN_MINUTES} league minutes - attack_p90 = 0.")
        return 0.0
    return ((top["npg"] + top["assists"]) * top["tier"]).sum() / (top["minutes"].sum() / 90)

def ucl_defenders(t):
    p = team_players(t)
    return int(((p["pos"].isin(["DF", "GK"])) & (p["ucl_min"] > 0)).sum()) if len(p) else np.nan

def h2h_gd(match_no):
    r = h2h[h2h["match_no"] == match_no]
    if r.empty or pd.isna(r["meetings"].iloc[0]):
        return np.nan
    n = float(r["meetings"].iloc[0])
    if n == 0:
        return 0.0
    return (float(r["a_goals"].iloc[0]) - float(r["b_goals"].iloc[0])) / n

# ------------------------------------------------------------------- team-level values per match
prior_cache = {}
def pre_tournament_form(t):
    if t not in prior_cache:
        prior_cache[t] = form(t, TOURNAMENT_START)
    return prior_cache[t]

def team_values(t, date):
    s, c = form(t, date)
    s0, c0 = pre_tournament_form(t)
    tx = tournament_xg(t, date)
    n = len(tx)
    xgd_mean = (tx["xgf"] - tx["xga"]).mean() if n else 0.0
    return dict(
        P=p_of(t) if t in P else np.nan,
        scored=s, conceded=c,
                   xg_diff_blend=blend(n, xgd_mean, 0.0),
        xg_for_blend=blend(n, tx["xgf"].mean() if n else 0.0, XG22),
        xga_blend=blend(n, tx["xga"].mean() if n else 0.0, XG22),
        ucl=ucl_minutes(t), attack=attack_output(t), p90=attack_p90(t), ucl_def=ucl_defenders(t),
        host=1 if t in HOSTS else 0,
    )

# ------------------------------------------------------------------- assemble datasets
rows21, rows22 = [], []
for r in m_filled.sort_values("match_no").itertuples():
    A, B = team_values(r.team_A, r.date), team_values(r.team_B, r.date)
    gA = pd.to_numeric(r.goals_A, errors="coerce")
    gB = pd.to_numeric(r.goals_B, errors="coerce")
    base = dict(match_no=r.match_no, date=r.date.date(), stage=r.stage)
    rows21.append(dict(**base, team_A=r.team_A, team_B=r.team_B,
        fifa_diff=A["P"] - B["P"],
        scored_diff=A["scored"] - B["scored"],
        conceded_diff=A["conceded"] - B["conceded"],
        xg_diff=A["xg_diff_blend"] - B["xg_diff_blend"],
        ucl_min_diff=A["ucl"] - B["ucl"],
        attack_out_diff=A["attack"] - B["attack"],
        host=A["host"] - B["host"],
        h2h_gd=h2h_gd(r.match_no),
        goal_diff=gA - gB, Set_80_20=r.Set_80_20, Set_70_30=r.Set_70_30))
    for side, T, O, team, opp, goals in (("A", A, B, r.team_A, r.team_B, gA),
                                         ("B", B, A, r.team_B, r.team_A, gB)):
        rows22.append(dict(**base, team=team, opponent=opp, side=side,
            fifa_diff=T["P"] - O["P"],
            score_form=T["scored"],
            opp_concede_form=O["conceded"],
            xg_for=T["xg_for_blend"],
            opp_xga=O["xga_blend"],
            attack_p90=T["p90"],
            host=T["host"],
            opp_ucl_def=O["ucl_def"],
            goals=goals, Set_80_20=r.Set_80_20, Set_70_30=r.Set_70_30))

lr21, lr22 = pd.DataFrame(rows21), pd.DataFrame(rows22)
lr21.to_csv("lr21.csv", index=False)
lr22.to_csv("lr22.csv", index=False)

print(f"lr21.csv: {len(lr21)} rows (target 104)   lr22.csv: {len(lr22)} rows (target 208)")
print(f"P_avg (average FIFA points of qualified teams) = {P_AVG:.1f}" if not np.isnan(P_AVG) else "P_avg: no FIFA points yet")
print("\nMissing values per variable (2.1):")
print(lr21.iloc[:, 5:13].isna().sum().to_string())
print("\nMissing values per variable (2.2):")
print(lr22.iloc[:, 6:14].isna().sum().to_string())
if warnings:
    uniq = list(dict.fromkeys(warnings))
    print(f"\n{len(uniq)} data warning(s) - fix these in the workbook:")
    for w in uniq[:60]:
        print("  -", w)
    if len(uniq) > 60:
        print(f"  ... and {len(uniq) - 60} more")
else:
    print("\nNo data warnings.")