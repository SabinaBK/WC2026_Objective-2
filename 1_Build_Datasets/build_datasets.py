import os
import re
import sys
import numpy as np
import pandas as pd
from openpyxl import load_workbook

# Folder of this script and the project folder
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
WORKBOOK = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "Data", "WC2026_data.xlsx")
TOURNAMENT_START = pd.Timestamp("2026-06-11")
HOSTS = {"USA", "Canada", "Mexico"}
BIG5 = {"premier league", "la liga", "bundesliga", "serie a", "ligue 1"}
N_FORM = 10          # number of recent competitive matches used for form
DECAY = 0.9          # weight of the match i places back is DECAY**i
K_BLEND = 3          # pre-tournament prior counts as K_BLEND matches when blending xG
MIN_MINUTES = 900    # minimum Big 5 league minutes for an attacker in attack_p90

warnings = []


def warn(msg):
    warnings.append(msg)


def norm(text):
    """Lower-case a header and drop spaces, underscores and symbols so headers can be matched loosely."""
    return re.sub(r"[^a-z0-9]", "", str(text).lower())


SHEETS = pd.ExcelFile(WORKBOOK).sheet_names


def read_sheet(names, columns, key):
    """Read a sheet and return only the needed columns, found by header name.
    names lists accepted sheet names; columns maps each internal name to accepted header spellings."""
    names = [names] if isinstance(names, str) else names
    name = next((s for s in SHEETS if norm(s) in {norm(n) for n in names}), None)
    if name is None:
        sys.exit(f"Sheet '{names[0]}' is missing from {WORKBOOK}.")
    raw = pd.read_excel(WORKBOOK, sheet_name=name, header=0)
    headers = {norm(c): c for c in raw.columns}
    out = pd.DataFrame()
    for col, aliases in columns.items():
        match = next((headers[norm(a)] for a in aliases if norm(a) in headers), None)
        if match is None:
            sys.exit(f"Sheet '{name}' has no column for '{col}' (accepted headers: {aliases}).")
        out[col] = raw[match]
    return out[out[key].notna()].copy()


# FBref short names mapped to the spelling used in FIFA_Points
NAME_FIX = {"Bosnia–Herz": "Bosnia and Herzegovina", "CAR": "Central African Republic",
            "No. Macedonia": "North Macedonia", "Rep. of Ireland": "Republic of Ireland",
            "Trinidad & Tobago": "Trinidad and Tobago", "UAE": "United Arab Emirates"}


def clean(s):
    """Trim a team name, unify the apostrophe in Côte d'Ivoire and apply NAME_FIX."""
    if not isinstance(s, str):
        return s
    s = s.strip().replace("\u2019", "'").replace("*", "").strip()
    return NAME_FIX.get(s, s)


def num(series):
    return pd.to_numeric(series, errors="coerce")


# ---------------------------------------------------------------- raw data
split = read_sheet("Split", {"match_no": ["match_no", "match no", "match_no."],
                             "Set_80_20": ["Set_80_20"], "Set_70_30": ["Set_70_30"]}, "match_no")
split["match_no"] = num(split["match_no"])

m = read_sheet("Matches", {"match_no": ["match_no", "match no"], "date": ["date"], "stage": ["stage", "round"],
                           "team_A": ["team_A", "home"], "team_B": ["team_B", "away"],
                           "goals_A": ["goals_A"], "goals_B": ["goals_B"]}, "match_no")
m["match_no"] = num(m["match_no"])
m["date"] = pd.to_datetime(m["date"], errors="coerce")
m["team_A"], m["team_B"] = m["team_A"].map(clean), m["team_B"].map(clean)
m = m.merge(split, on="match_no", how="left").sort_values("match_no")

fp = read_sheet("FIFA_Points", {"team": ["team"], "points": ["FIFA_Points", "points", "fifa points"]}, "team")
fp["team"] = fp["team"].map(clean)
P = dict(zip(fp["team"], num(fp["points"])))
WC_TEAMS = sorted(set(m["team_A"]) | set(m["team_B"]))
# Average points of the 48 World Cup teams, used to scale opponent strength
P_AVG = np.nanmean([P.get(t, np.nan) for t in WC_TEAMS])

rm = read_sheet("Recent_Matches", {"team": ["team"], "date": ["date"], "opponent": ["opponent"],
                                   "competition": ["competition"],
                                   "gf": ["goals_for", "gf"], "ga": ["goals_against", "ga"]}, "date")
rm["date"] = pd.to_datetime(rm["date"], errors="coerce")
rm["team"], rm["opponent"] = rm["team"].map(clean), rm["opponent"].map(clean)
# Only competitive matches played before the tournament; World Cup matches are added from Matches
friendly = rm["competition"].astype(str).str.contains("friendl", case=False)
if friendly.any():
    warn(f"Recent_Matches: {int(friendly.sum())} friendly row(s) removed.")
rm = rm[~friendly & (rm["date"] < TOURNAMENT_START)]

ms = read_sheet("Match_Stats", {"match_no": ["match_no", "match no"], "xg": ["xg"]}, "match_no")
ms["match_no"] = num(ms["match_no"])
# Within each match the first row is Team A and the second row is Team B
ms["side"] = ms.groupby("match_no").cumcount().map({0: "A", 1: "B"})

pl = read_sheet(["Players", "Player"], {"team": ["team"], "pos": ["position", "pos"], "league": ["league"],
                            "minutes": ["league_minutes", "minutes"],
                            "npg": ["non_penalty_goals", "npg"], "assists": ["assists"],
                            "ucl_min": ["champions_league_minutes", "ucl_minutes"]}, "team")
pl["team"] = pl["team"].map(clean)
pl["pos"] = pl["pos"].astype(str).str.upper().str.strip()
for c in ("minutes", "npg", "assists", "ucl_min"):
    pl[c] = num(pl[c]).fillna(0)


def is_big5(league):
    """True if any league in the cell (e.g. 'eng Premier League / fr Ligue 1') is a Big 5 league."""
    if not isinstance(league, str):
        return False
    for part in league.split("/"):
        name = re.sub(r"^[a-z]{2,3}\s+", "", part.strip())   # remove FBref country prefix such as 'eng '
        if name.lower() in BIG5:
            return True
    return False


# Player output variables count Big 5 league football only; other players contribute zero
pl["big5"] = pl["league"].map(is_big5)
pl.loc[~pl["big5"], ["minutes", "npg", "assists"]] = 0

h2h = read_sheet("H2H", {"match_no": ["match_no", "match no", "fixture #", "fixture"],
                         "meetings": ["meetings", "previous meetings"],
                         "a_goals": ["A_goals", "team a goals"], "b_goals": ["B_goals", "team b goals"]}, "match_no")
h2h["match_no"] = num(h2h["match_no"])

wb = load_workbook(WORKBOOK, data_only=True, read_only=True)
XG22 = wb["Constants"]["B3"].value if "Constants" in wb.sheetnames else None
if XG22 is None:
    sys.exit("Constants!B3 (2022 World Cup average per team per match) is empty.")
XG22 = float(XG22)

# ---------------------------------------------------------------- data checks
for t in WC_TEAMS:
    if pd.isna(P.get(t)):
        warn(f"FIFA_Points: no points for '{t}'.")
    if (pl["team"] == t).sum() == 0:
        warn(f"Players: no players for '{t}'.")
# Teams with fewer than N_FORM recent matches use the matches they have
counts = rm["team"].value_counts().reindex(WC_TEAMS, fill_value=0)
short = counts[counts < N_FORM]
if len(short):
    warn(f"Recent_Matches: {len(short)} team(s) have fewer than {N_FORM} matches (form uses what is available): "
         + ", ".join(f"{t} {n}" for t, n in short.sort_values().items()))
for opp in sorted(set(rm["opponent"].dropna())):
    if pd.isna(P.get(opp)):
        warn(f"FIFA_Points: no points for recent opponent '{opp}' (average points used).")


def points(team):
    """FIFA points of a team; average points if the team is missing."""
    v = P.get(team, np.nan)
    return P_AVG if pd.isna(v) else v


# ---------------------------------------------------------------- recent form
# One row per team per match: pre-tournament competitive matches plus World Cup 2026 matches
hist = [dict(team=r.team, date=r.date, gf=r.gf, ga=r.ga, opp=r.opponent) for r in rm.itertuples()]
for r in m.dropna(subset=["goals_A", "goals_B"]).itertuples():
    hist.append(dict(team=r.team_A, date=r.date, gf=r.goals_A, ga=r.goals_B, opp=r.team_B))
    hist.append(dict(team=r.team_B, date=r.date, gf=r.goals_B, ga=r.goals_A, opp=r.team_A))
hist = pd.DataFrame(hist)
hist["gf"], hist["ga"] = num(hist["gf"]), num(hist["ga"])
hist["p_opp"] = hist["opp"].map(points)


def raw_form(team, before):
    """Opponent-adjusted goals scored and conceded per match over the last N_FORM matches before a date.
    Goals against strong opponents count more; goals conceded to strong opponents count less."""
    h = hist[(hist["team"] == team) & (hist["date"] < before)].dropna(subset=["gf", "ga"])
    h = h.sort_values("date", ascending=False).head(N_FORM)
    if h.empty:
        return np.nan, np.nan
    w = DECAY ** np.arange(len(h))                      # most recent match has weight 1
    gf, ga, po = (h[c].to_numpy(float) for c in ("gf", "ga", "p_opp"))
    scored = np.sum(w * gf * po / P_AVG) / w.sum()
    conceded = np.sum(w * ga * P_AVG / po) / w.sum()
    return scored, conceded


# Average pre-tournament form of the teams that have recent matches; used for a team with none
pre = pd.DataFrame([raw_form(t, TOURNAMENT_START) for t in WC_TEAMS], index=WC_TEAMS,
                   columns=["scored", "conceded"])
DEFAULT_FORM = (pre["scored"].mean(), pre["conceded"].mean())
for t in pre.index[pre["scored"].isna()]:
    warn(f"Recent_Matches: '{t}' has no matches before the tournament; average form used for its first match.")


def form(team, before):
    scored, conceded = raw_form(team, before)
    return DEFAULT_FORM if np.isnan(scored) else (scored, conceded)


def raw_rates(team, before):
    """Weighted points per match (3 win, 1 draw, 0 loss) and share of matches without scoring,
    over the same last N_FORM matches used for form."""
    h = hist[(hist["team"] == team) & (hist["date"] < before)].dropna(subset=["gf", "ga"])
    h = h.sort_values("date", ascending=False).head(N_FORM)
    if h.empty:
        return np.nan, np.nan
    w = DECAY ** np.arange(len(h))
    gf, ga = h["gf"].to_numpy(float), h["ga"].to_numpy(float)
    pts = np.where(gf > ga, 3, np.where(gf == ga, 1, 0))
    return np.sum(w * pts) / w.sum(), np.sum(w * (gf == 0)) / w.sum()


# Average pre-tournament rates, used for a team with no recent matches
pre_rates = pd.DataFrame([raw_rates(t, TOURNAMENT_START) for t in WC_TEAMS], columns=["ppm", "fts"])
DEFAULT_RATES = (pre_rates["ppm"].mean(), pre_rates["fts"].mean())


def rates(team, before):
    ppm, fts = raw_rates(team, before)
    return DEFAULT_RATES if np.isnan(ppm) else (ppm, fts)


# ---------------------------------------------------------------- in-tournament xG
xg = ms.merge(m[["match_no", "date", "team_A", "team_B"]], on="match_no", how="inner")
xg["team"] = np.where(xg["side"] == "A", xg["team_A"], xg["team_B"])
xg["xg"] = num(xg["xg"])
opp_xg = xg[["match_no", "side", "xg"]].assign(side=xg["side"].map({"A": "B", "B": "A"}))
xg = xg.merge(opp_xg.rename(columns={"xg": "xga"}), on=["match_no", "side"], how="left")
xg = xg.rename(columns={"xg": "xgf"})[["team", "date", "match_no", "xgf", "xga"]]
missing_xg = sorted(int(x) for x in xg.loc[xg["xgf"].isna(), "match_no"].unique())
if missing_xg:
    warn(f"Match_Stats: xG missing for match(es) {missing_xg}; excluded from teams' xG averages.")


def tournament_xg(team, before):
    """The team's earlier World Cup 2026 matches with xG recorded for both sides."""
    h = xg[(xg["team"] == team) & (xg["date"] < before)]
    return h.dropna(subset=["xgf", "xga"])


def blend(n, mean, prior):
    """Weighted average of the tournament mean (n matches) and a prior worth K_BLEND matches."""
    return prior if n == 0 else (n * mean + K_BLEND * prior) / (n + K_BLEND)


# ---------------------------------------------------------------- player variables
def squad(team):
    return pl[pl["team"] == team]


def ucl_minutes(team):
    return squad(team)["ucl_min"].sum()


def attack_output(team):
    """Total Big 5 league non-penalty goals plus assists across the squad."""
    p = squad(team)
    return (p["npg"] + p["assists"]).sum()


def attack_p90(team):
    """Non-penalty goals plus assists per 90 minutes of the three most-used attackers
    (forwards first, then midfielders) with at least MIN_MINUTES Big 5 league minutes."""
    p = squad(team)
    p = p[p["minutes"] >= MIN_MINUTES]
    top = pd.concat([p[p["pos"] == "FW"].sort_values("minutes", ascending=False),
                     p[p["pos"] == "MF"].sort_values("minutes", ascending=False)]).head(3)
    if top.empty:
        return 0.0
    return (top["npg"] + top["assists"]).sum() / (top["minutes"].sum() / 90)


def ucl_defenders(team):
    """Number of defenders and goalkeepers who played in the 2025-26 Champions League."""
    p = squad(team)
    return int((p["pos"].isin(["DF", "GK"]) & (p["ucl_min"] > 0)).sum())


def h2h_gd(match_no):
    """Average goal difference (Team A minus Team B) in previous meetings; 0 if they never met."""
    r = h2h[h2h["match_no"] == match_no]
    if r.empty or pd.isna(r["meetings"].iloc[0]):
        return np.nan
    n = float(r["meetings"].iloc[0])
    return 0.0 if n == 0 else (float(r["a_goals"].iloc[0]) - float(r["b_goals"].iloc[0])) / n


# ---------------------------------------------------------------- team values at a given match date
def team_values(team, date):
    scored, conceded = form(team, date)
    ppm, fts = rates(team, date)
    tx = tournament_xg(team, date)
    n = len(tx)
    return dict(
        P=P.get(team, np.nan),
        scored=scored,
        conceded=conceded,
        ppm=ppm,
        fts=fts,
        xg_diff=blend(n, (tx["xgf"] - tx["xga"]).mean() if n else 0.0, 0.0),
        xg_for=blend(n, tx["xgf"].mean() if n else 0.0, XG22),
        xga=blend(n, tx["xga"].mean() if n else 0.0, XG22),
        ucl=ucl_minutes(team),
        attack=attack_output(team),
        p90=attack_p90(team),
        ucl_def=ucl_defenders(team),
        host=int(team in HOSTS),
    )


# ---------------------------------------------------------------- assemble datasets
rows21, rows22 = [], []
for r in m.itertuples():
    A, B = team_values(r.team_A, r.date), team_values(r.team_B, r.date)
    gA, gB = pd.to_numeric(r.goals_A, errors="coerce"), pd.to_numeric(r.goals_B, errors="coerce")
    base = dict(match_no=int(r.match_no), date=r.date.date(), stage=r.stage)

    # Linear Regression 2.1: one row per match
    rows21.append(dict(**base, team_A=r.team_A, team_B=r.team_B,
                       fifa_diff=A["P"] - B["P"],
                       scored_diff=A["scored"] - B["scored"],
                       conceded_diff=A["conceded"] - B["conceded"],
                       xg_diff=A["xg_diff"] - B["xg_diff"],
                       ucl_min_diff=A["ucl"] - B["ucl"],
                       attack_out_diff=A["attack"] - B["attack"],
                       host=A["host"] - B["host"],
                       h2h_gd=h2h_gd(r.match_no),
                       ppm_diff=A["ppm"] - B["ppm"],          # candidate replacement for ucl_min_diff
                       goal_diff=gA - gB,
                       Set_80_20=r.Set_80_20, Set_70_30=r.Set_70_30))

    # Linear Regression 2.2: two rows per match, Team A's view then Team B's view
    for side, T, O, team, opp, goals in (("A", A, B, r.team_A, r.team_B, gA),
                                         ("B", B, A, r.team_B, r.team_A, gB)):
        rows22.append(dict(**base, team=team, opponent=opp, side=side,
                           fifa_diff=T["P"] - O["P"],
                           score_form=T["scored"],
                           opp_concede_form=O["conceded"],
                           xg_for=T["xg_for"],
                           opp_xga=O["xga"],
                           attack_p90=T["p90"],
                           host=T["host"],
                           opp_ucl_def=O["ucl_def"],
                           fail_to_score=T["fts"],                # candidate replacement for opp_ucl_def
                           goals=goals,
                           Set_80_20=r.Set_80_20, Set_70_30=r.Set_70_30))

lr21, lr22 = pd.DataFrame(rows21), pd.DataFrame(rows22)
lr21.to_csv(os.path.join(HERE, "LR2.1_dataset.csv"), index=False)
lr22.to_csv(os.path.join(HERE, "LR2.2_dataset.csv"), index=False)

VARS21 = ["fifa_diff", "scored_diff", "conceded_diff", "xg_diff",
          "ucl_min_diff", "attack_out_diff", "host", "h2h_gd", "ppm_diff"]
VARS22 = ["fifa_diff", "score_form", "opp_concede_form", "xg_for",
          "opp_xga", "attack_p90", "host", "opp_ucl_def", "fail_to_score"]

print(f"LR2.1_dataset.csv: {len(lr21)} rows    LR2.2_dataset.csv: {len(lr22)} rows")
print(f"Average FIFA points of the 48 teams: {P_AVG:.1f}")
print(f"Big 5 league players: {int(pl['big5'].sum())} of {len(pl)}")
print("\nMissing values (2.1):\n" + lr21[VARS21].isna().sum().to_string())
print("\nMissing values (2.2):\n" + lr22[VARS22].isna().sum().to_string())
if warnings:
    print(f"\n{len(warnings)} warning(s):")
    for w in dict.fromkeys(warnings):
        print("  -", w)